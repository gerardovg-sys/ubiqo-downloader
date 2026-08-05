# -*- coding: utf-8 -*-
"""
ANÁLISIS EXHAUSTIVO BITÁCORA SMARTCORP
Hoja: Bitácora (3a hoja)
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

import openpyxl
import pandas as pd
from collections import Counter
from datetime import datetime, timedelta
import json

FILE = r"c:\Users\MARCO\Desktop\HOLDING_IA\BITACORA_SMARTCORP\scratch\BD Proyectos_ejemplo.xlsx"
SHEET = "Bitácora"

print("=" * 70)
print("PASO 1 — EXPLORACIÓN DE ESTRUCTURA")
print("=" * 70)

wb = openpyxl.load_workbook(FILE, read_only=True, data_only=True)
print(f"\nHojas disponibles ({len(wb.sheetnames)}): {wb.sheetnames}")

ws = wb[SHEET]
all_rows = list(ws.iter_rows(values_only=True))
wb.close()

header = list(all_rows[0])
data = [r for r in all_rows[1:] if any(v is not None for v in r)]
df = pd.DataFrame(data, columns=header)

print(f"\nHoja analizada: '{SHEET}'")
print(f"Total columnas: {len(header)}")
print(f"Total filas de datos: {len(df)}")
print(f"\nColumnas completas:")
for i, h in enumerate(header):
    print(f"  [{i:>2}] {repr(h)}")

print("\n--- PRIMERAS 5 FILAS ---")
for row_num, row in enumerate(data[:5], start=2):
    print(f"\n  Fila Excel {row_num}:")
    for i, (h, v) in enumerate(zip(header, row)):
        if v is not None:
            print(f"    [{i}] {h}: {repr(v)}")


print("\n\n" + "=" * 70)
print("PASO 2 — ANÁLISIS DE COLUMNAS CLAVE")
print("=" * 70)

# === FECHA ===
print("\n--- FECHA ---")
fechas_raw = df['FECHA'].dropna()
parsed_dates = []
for v in fechas_raw:
    if isinstance(v, datetime):
        parsed_dates.append(v.date())
    else:
        try:
            parsed_dates.append(pd.to_datetime(v).date())
        except:
            pass

print(f"  Con fecha: {len(parsed_dates)}/{len(df)}")
if parsed_dates:
    d_min = min(parsed_dates)
    d_max = max(parsed_dates)
    d_uniq = set(parsed_dates)
    print(f"  Rango: {d_min} → {d_max}")
    print(f"  Días únicos: {len(d_uniq)} (de {(d_max - d_min).days + 1} días en rango)")
    fmt_det = "datetime nativo Excel" if isinstance(fechas_raw.iloc[0], datetime) else "string"
    print(f"  Formato: {fmt_det}")
    mc = Counter(d.strftime("%Y-%m") for d in parsed_dates)
    print(f"\n  Distribución mensual:")
    for m, c in sorted(mc.items()):
        bar = "#" * (c // 5)
        print(f"    {m}: {c:>4}  {bar}")

# === UNIDAD ===
print("\n--- UNIDAD ---")
u_series = df['UNIDAD'].dropna().astype(str).str.strip()
u_series = u_series[u_series != '']
cu = Counter(u_series)
print(f"  Con unidad: {len(u_series)}/{len(df)}")
print(f"  Sin unidad: {len(df) - len(u_series)}/{len(df)} ({100*(len(df)-len(u_series))/len(df):.1f}%)")
print(f"  Unidades únicas: {len(cu)}")
print(f"\n  Todas las unidades (por frecuencia):")
for k, v in sorted(cu.items(), key=lambda x: -x[1]):
    print(f"    {k:<35} {v:>4} registros")

# === PROYECTO ===
print("\n--- PROYECTO ---")
p_series = df['PROYECTO'].dropna().astype(str).str.strip()
cp = Counter(p_series)
print(f"  Con proyecto: {len(p_series)}/{len(df)}")
print(f"  Proyectos únicos: {len(cp)}")
print(f"\n  Top 25 proyectos más frecuentes:")
for i, (k, v) in enumerate(cp.most_common(25), 1):
    print(f"    {i:>2}. {k[:60]:<60} {v:>4}")

# === NOMBRE ===
print("\n--- NOMBRE (TÉCNICO) ---")
n_series = df['NOMBRE'].dropna().astype(str).str.strip()
cn = Counter(n_series)
print(f"  Con nombre: {len(n_series)}/{len(df)}")
print(f"  Técnicos únicos: {len(cn)}")
print(f"\n  Todos los técnicos:")
for k, v in sorted(cn.items(), key=lambda x: -x[1]):
    print(f"    {k:<45} {v:>4} registros")

# Técnicos por día
df['_fecha_str'] = df['FECHA'].apply(
    lambda x: x.date().isoformat() if isinstance(x, datetime) else str(x))
df_cn = df.dropna(subset=['NOMBRE'])
tec_dia = df_cn.groupby('_fecha_str')['NOMBRE'].nunique()
print(f"\n  Promedio técnicos únicos/día: {tec_dia.mean():.2f}")
print(f"  Máximo técnicos en un día: {tec_dia.max()}")

# === ASUNTO ===
print("\n--- ASUNTO ---")
a_series = df['ASUNTO'].dropna().astype(str).str.strip()
ca = Counter(a_series)
print(f"  Con asunto: {len(a_series)}/{len(df)}")
print(f"  Valores únicos: {len(ca)}")
print(f"\n  Todos los valores de asunto:")
for k, v in sorted(ca.items(), key=lambda x: -x[1]):
    pct = 100 * v / len(df)
    print(f"    {repr(k):<50} {v:>5}  ({pct:.1f}%)")

# Clasificar
inst_kws = ["instalaci", "proyecto"]
def cat_asunto(a):
    a = str(a).lower()
    if any(k in a for k in inst_kws): return "INSTALACION"
    elif "retro" in a: return "RETROALIMENTACION"
    elif "oficina" in a: return "OFICINA"
    elif "capacitaci" in a: return "CAPACITACION"
    else: return "OTRO"

df['_cat'] = df['ASUNTO'].apply(lambda x: cat_asunto(x) if pd.notna(x) else "SIN_ASUNTO")
cat_counts = Counter(df['_cat'])
print(f"\n  Clasificación de asuntos:")
for k, v in sorted(cat_counts.items(), key=lambda x: -x[1]):
    print(f"    {k:<20} {v:>5}  ({100*v/len(df):.1f}%)")

# === GPS COLUMNS ===
print("\n--- COLUMNAS GPS (% LLENADO) ---")
gps_cols = [
    'HORA DE SALIDA', 'HORA DE ENTRADA', 'TIEMPO RECORRIDO',
    'TIEMPO DE PARADAS', 'PARADAS', 'REGRESOS', 'KM',
    'HORA SAL PROY', 'HORA LLEG PROY', 'ASISTENCIA', 'HORAS EXTRA'
]
total = len(df)
for col in gps_cols:
    if col in df.columns:
        filled = df[col].notna().sum()
        pct = 100 * filled / total
        sample = [repr(v) for v in df[col].dropna().head(3).tolist()]
        tipo = type(df[col].dropna().iloc[0]).__name__ if filled > 0 else "N/A"
        print(f"  {col:<25} {filled:>4}/{total}  ({pct:>5.1f}%)  tipo={tipo}")
        if sample:
            print(f"    Muestra: {', '.join(sample)}")
    else:
        print(f"  {col:<25} -> NO ENCONTRADA")

# === OTRAS COLUMNAS ===
print("\n--- OTRAS COLUMNAS RELEVANTES ---")
otros = ['q','SAP','Rol','DE','A','CALCULO HORAS','REPORTE ENV.','JUSTIFICACION','NOTA','REV','OBSERVACIONES']
for col in otros:
    if col in df.columns:
        filled = df[col].notna().sum()
        pct = 100 * filled / total
        sample_vals = df[col].dropna().unique()[:5].tolist()
        print(f"  {col:<20} {filled:>4}/{total}  ({pct:>5.1f}%)  muestra: {sample_vals}")


print("\n\n" + "=" * 70)
print("PASO 3 — PATRONES OPERATIVOS")
print("=" * 70)

# PAT1: proyectos por unidad-día
print("\n--- PAT1: Proyectos por Unidad por Día ---")
df_up = df[df['UNIDAD'].notna() & df['PROYECTO'].notna() &
           (df['UNIDAD'].astype(str).str.strip() != '')].copy()
df_up['_u'] = df_up['UNIDAD'].astype(str).str.strip()
df_up['_p'] = df_up['PROYECTO'].astype(str).str.strip()

gproy = df_up.groupby(['_u', '_fecha_str'])['_p'].nunique()
dist = Counter(gproy.values)
print(f"  Distribución (Unidad, Día) → N proyectos únicos:")
for num, freq in sorted(dist.items()):
    print(f"    {num} proyecto(s): {freq} casos")
print(f"\n  Promedio: {gproy.mean():.2f}  |  Máximo: {gproy.max()}")
print(f"  Casos con 3+ proyectos: {(gproy >= 3).sum()}")

multi_proy = gproy[gproy >= 2]
print(f"\n  Ejemplos días con 2+ proyectos (primeros 15):")
for i, ((u, f), n) in enumerate(multi_proy.items()):
    if i >= 15: break
    mask = (df_up['_u'] == u) & (df_up['_fecha_str'] == f)
    projs = df_up[mask]['_p'].unique().tolist()
    noms = df_up[mask]['NOMBRE'].dropna().unique().tolist()
    print(f"    {u} | {f} → {projs}")
    print(f"      Personas: {noms}")

# PAT2: Múltiples personas en misma unidad-día
print("\n--- PAT2: Múltiples Personas en Misma Unidad/Día ---")
gpers = df_up.groupby(['_u', '_fecha_str'])['NOMBRE'].nunique()
multi_pers = gpers[gpers >= 2]
print(f"  Casos con 2+ personas compartiendo unidad: {len(multi_pers)}")

if len(multi_pers) > 0:
    print(f"\n  Ejemplos (primeros 10):")
    for i, ((u, f), n) in enumerate(multi_pers.items()):
        if i >= 10: break
        mask = (df_up['_u'] == u) & (df_up['_fecha_str'] == f)
        personas = df_up[mask]['NOMBRE'].dropna().unique().tolist()
        proyectos = df_up[mask]['_p'].unique().tolist()
        print(f"    {u} | {f}")
        print(f"      Personas ({n}): {personas}")
        print(f"      Proyectos: {proyectos}")
    
    mismo = 0
    diff = 0
    for (u, f), n in multi_pers.items():
        mask = (df_up['_u'] == u) & (df_up['_fecha_str'] == f)
        n_proy = df_up[mask]['_p'].nunique()
        if n_proy == 1: mismo += 1
        else: diff += 1
    print(f"\n  → MISMO proyecto: {mismo}  |  DIFERENTE proyecto: {diff}")

# PAT3: Asuntos mixtos
print("\n--- PAT3: Días Mixtos (Instalación + Otro Asunto) por Persona ---")
df_cat = df.dropna(subset=['NOMBRE', '_fecha_str', 'ASUNTO'])
g3 = df_cat.groupby(['NOMBRE', '_fecha_str'])['_cat'].apply(set)
mix_inst_retro = g3[g3.apply(lambda s: "INSTALACION" in s and "RETROALIMENTACION" in s)]
mix_inst_ofic = g3[g3.apply(lambda s: "INSTALACION" in s and "OFICINA" in s)]
mix_any_inst = g3[g3.apply(lambda s: "INSTALACION" in s and len(s) > 1)]
print(f"  Persona-días con INSTALACION + RETROALIMENTACION: {len(mix_inst_retro)}")
print(f"  Persona-días con INSTALACION + OFICINA: {len(mix_inst_ofic)}")
print(f"  Persona-días con INSTALACION + CUALQUIER OTRO: {len(mix_any_inst)}")

if len(mix_any_inst) > 0:
    print(f"\n  Ejemplos mixtos (primeros 8):")
    for i, ((nom, fecha), cats) in enumerate(mix_any_inst.items()):
        if i >= 8: break
        asuntos_detalle = df_cat[(df_cat['NOMBRE'] == nom) & (df_cat['_fecha_str'] == fecha)]['ASUNTO'].unique().tolist()
        print(f"    {nom} | {fecha} → {asuntos_detalle}")

# PAT4: Secuencias
print("\n--- PAT4: Distribución de Secuencias ---")
print(f"  Días con 1 proyecto: {dist.get(1,0)}")
print(f"  Días con 2 proyectos: {dist.get(2,0)}")
print(f"  Días con 3 proyectos: {dist.get(3,0)}")
print(f"  Días con 4+ proyectos: {sum(v for k,v in dist.items() if k>=4)}")

# PAT5: Rango horario
print("\n--- PAT5: Rango Horario ---")
hora_cols = [('HORA DE SALIDA','salida'), ('HORA DE ENTRADA','entrada'),
             ('HORA SAL PROY','sal_proy'), ('HORA LLEG PROY','lleg_proy')]
for col, label in hora_cols:
    if col not in df.columns: continue
    vals = df[col].dropna()
    if len(vals) == 0:
        print(f"  {col}: VACÍA")
        continue
    print(f"\n  {col} ({len(vals)} registros):")
    print(f"    Tipo: {type(vals.iloc[0]).__name__}")
    print(f"    Primeros 5: {vals.head(5).tolist()}")
    horas = []
    for v in vals:
        if isinstance(v, datetime):
            horas.append(v.hour + v.minute/60.0)
        elif hasattr(v, 'hour'):  # datetime.time
            horas.append(v.hour + v.minute/60.0)
        elif isinstance(v, (int, float)) and 0 < v < 1:
            horas.append(v * 24)
    if horas:
        hmin = min(horas)
        hmax = max(horas)
        hmean = sum(horas)/len(horas)
        print(f"    Hora min: {int(hmin):02d}:{int((hmin%1)*60):02d}  max: {int(hmax):02d}:{int((hmax%1)*60):02d}  prom: {int(hmean):02d}:{int((hmean%1)*60):02d}")
        nocturnos = [h for h in horas if h < 5 or h > 22]
        print(f"    Registros nocturnos (<5h o >22h): {len(nocturnos)}")


print("\n\n" + "=" * 70)
print("PASO 4 — ANOMALÍAS Y CASOS ESPECIALES")
print("=" * 70)

# AN1
print("\n--- AN1: Filas sin UNIDAD ---")
sin_u = df[df['UNIDAD'].isna() | (df['UNIDAD'].astype(str).str.strip() == '')]
print(f"  Sin UNIDAD: {len(sin_u)}/{len(df)} ({100*len(sin_u)/len(df):.1f}%)")
ca_sin_u = Counter(sin_u['ASUNTO'].dropna().astype(str).str.strip())
print(f"  Sus asuntos:")
for k, v in ca_sin_u.most_common(10):
    print(f"    '{k}': {v}")

# AN2
print("\n--- AN2: UNIDAD llena, PROYECTO vacío ---")
tu = df[df['UNIDAD'].notna() & (df['UNIDAD'].astype(str).str.strip() != '')]
sp = tu[tu['PROYECTO'].isna() | (tu['PROYECTO'].astype(str).str.strip() == '')]
print(f"  Con UNIDAD sin PROYECTO: {len(sp)}/{len(tu)} ({100*len(sp)/len(tu):.1f}%)")
ca_sp = Counter(sp['ASUNTO'].dropna().astype(str).str.strip())
for k, v in ca_sp.most_common(10):
    print(f"    '{k}': {v}")

# AN3: Proyectos especiales MTTO
print("\n--- AN3: Proyectos tipo MTTO PERSONAL ---")
kws_bl = ["MTTO PERSONAL", "INT QRO MTTO", "MTTO"]
for kw in kws_bl:
    mask = df['PROYECTO'].astype(str).str.upper().str.contains(kw, na=False)
    bl = df[mask]
    print(f"\n  Keyword '{kw}': {len(bl)} filas")
    if len(bl) > 0:
        proyectos_exactos = bl['PROYECTO'].dropna().unique().tolist()
        for p in proyectos_exactos[:10]:
            print(f"    Proyecto: '{p}'")
        asuntos_bl = Counter(bl['ASUNTO'].dropna().astype(str).str.strip())
        for a, c in asuntos_bl.items():
            print(f"    Asunto: '{a}' → {c}")

# AN4: GPS fill rate por proyecto
print("\n--- AN4: Top proyectos vs GPS (HORA DE SALIDA) ---")
df['_has_gps'] = df['HORA DE SALIDA'].apply(
    lambda x: x is not None and str(x).strip() not in ['', 'None', 'nan'])
top20 = [k for k, _ in Counter(df['PROYECTO'].dropna().astype(str).str.strip()).most_common(20)]
print(f"  {'Proyecto':<60} {'GPS_lleno':>10} {'Total':>7} {'%':>5}")
for proy in top20:
    mask_p = df['PROYECTO'].astype(str).str.strip() == proy
    tot_p = mask_p.sum()
    gps_p = df[mask_p]['_has_gps'].sum()
    pct = 100 * gps_p / tot_p if tot_p > 0 else 0
    flag = "✓" if pct > 50 else "✗"
    print(f"  {flag} {proy[:60]:<60} {gps_p:>10} {tot_p:>7}  {pct:>4.0f}%")

# AN5: Unidad llena, GPS vacío
print("\n--- AN5: UNIDAD llena pero HORA DE SALIDA vacía ---")
tu2 = df[df['UNIDAD'].notna() & (df['UNIDAD'].astype(str).str.strip() != '')]
sin_gps = tu2[tu2['HORA DE SALIDA'].apply(
    lambda x: x is None or str(x).strip() in ['', 'None', 'nan'])]
print(f"  Con UNIDAD pero sin HORA DE SALIDA: {len(sin_gps)}/{len(tu2)} ({100*len(sin_gps)/len(tu2):.1f}%)")
print(f"  Asuntos dominantes en esas filas:")
ca_sg = Counter(sin_gps['ASUNTO'].dropna().astype(str).str.strip())
for k, v in ca_sg.most_common(10):
    pct = 100 * v / len(sin_gps)
    print(f"    '{k}': {v} ({pct:.1f}%)")

# AN6: Columna ASISTENCIA
print("\n--- AN6: Valores de ASISTENCIA ---")
if 'ASISTENCIA' in df.columns:
    ca_asist = Counter(df['ASISTENCIA'].dropna().astype(str).str.strip())
    for k, v in ca_asist.most_common():
        print(f"    '{k}': {v}")

# AN7: Columna Rol
print("\n--- AN7: Valores de ROL ---")
if 'Rol' in df.columns:
    ca_rol = Counter(df['Rol'].dropna().astype(str).str.strip())
    for k, v in ca_rol.most_common():
        print(f"    '{k}': {v}")

print("\n\n" + "=" * 70)
print("RESUMEN EJECUTIVO")
print("=" * 70)
print(f"""
ARCHIVO: BD Proyectos_ejemplo.xlsx
HOJA: 'Bitácora'
TOTAL FILAS: {len(df)}
TOTAL COLUMNAS: {len(header)}
COLUMNAS: {header}

MESES CUBIERTOS: {sorted(set(d.strftime('%Y-%m') for d in parsed_dates)) if parsed_dates else 'N/A'}
DÍAS ÚNICOS: {len(set(parsed_dates)) if parsed_dates else 0}

GPS COLUMNS IDENTIFICADAS (en bitácora):
  - HORA DE SALIDA    (col 15)
  - HORA DE ENTRADA   (col 16)
  - TIEMPO RECORRIDO  (col 17)
  - TIEMPO DE PARADAS (col 18)
  - PARADAS           (col 19)
  - REGRESOS          (col 20)
  - KM                (col 22)
  - HORA SAL PROY     (col 24)
  - HORA LLEG PROY    (col 25)
""")

# Guardar column map
col_map = {
    "SHEET": SHEET,
    "FECHA": "FECHA", "UNIDAD": "UNIDAD", "PROYECTO": "PROYECTO",
    "NOMBRE": "NOMBRE", "ASUNTO": "ASUNTO", "ROL": "Rol",
    "HORA_SALIDA": "HORA DE SALIDA", "HORA_ENTRADA": "HORA DE ENTRADA",
    "TIEMPO_RECORRIDO": "TIEMPO RECORRIDO", "TIEMPO_PARADAS": "TIEMPO DE PARADAS",
    "PARADAS": "PARADAS", "REGRESOS": "REGRESOS", "KM": "KM",
    "HORA_SAL_PROY": "HORA SAL PROY", "HORA_LLEG_PROY": "HORA LLEG PROY",
    "ASISTENCIA": "ASISTENCIA", "HORAS_EXTRA": "HORAS EXTRA"
}
with open(r"c:\Users\MARCO\Desktop\HOLDING_IA\BITACORA_SMARTCORP\scratch\column_map.json","w",encoding="utf-8") as f:
    json.dump(col_map, f, ensure_ascii=False, indent=2)
print("column_map.json guardado.")
print("\n✅ ANÁLISIS COMPLETADO.")
