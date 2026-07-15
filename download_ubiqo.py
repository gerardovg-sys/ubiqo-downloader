import os
import sys
import json
import time
import base64
import re
import shutil
from datetime import datetime, timedelta, timezone

# Auto-instalar dependencias de Python si no están
try:
    import requests
except ImportError:
    print("Instalando la librería 'requests'...")
    os.system("pip install requests")
    import requests

try:
    from playwright.sync_api import sync_playwright
except ImportError:
    print("Instalando la librería Playwright para automatización...")
    os.system("pip install playwright")
    os.system("python -m playwright install chromium")
    from playwright.sync_api import sync_playwright

CONFIG_FILE = "ubiqo_config.json"

def cargar_config():
    if os.path.exists(CONFIG_FILE):
        with open(CONFIG_FILE, 'r', encoding='utf-8-sig') as f:
            return json.load(f)
    return {}

def guardar_config(config):
    with open(CONFIG_FILE, 'w', encoding='utf-8-sig') as f:
        json.dump(config, f, indent=2, ensure_ascii=False)

def obtener_credenciales():
    # Intentar obtener de variables de entorno (para ejecución en la nube / GitHub Actions)
    env_user = os.environ.get("UBIQO_USER")
    env_pass = os.environ.get("UBIQO_PASS")
    env_web_app = os.environ.get("WEB_APP_URL")
    env_token = os.environ.get("WEB_APP_TOKEN", "SMARTCORP_UBIQO_SECURE_TOKEN_2026")
    
    if env_user and env_pass:
        print("Cargando credenciales desde variables de entorno (Nube)...")
        return env_user, env_pass, env_web_app, env_token

    config = cargar_config()
    
    web_app_url = config.get("web_app_url")
    token = config.get("token", "SMARTCORP_UBIQO_SECURE_TOKEN_2026")
    user = config.get("user")
    password = config.get("pass")
    
    if not user or not password:
        print("\n=== Configuración Inicial de Ubiqo ===")
        user = input("Introduce tu Usuario de Ubiqo: ").strip()
        password = input("Introduce tu Contraseña de Ubiqo: ").strip()
        
        config["user"] = user
        config["pass"] = password
        guardar_config(config)
        
    return user, password, web_app_url, token

def subir_a_google_drive(file_path, web_app_url, token):
    filename = os.path.basename(file_path)
    print(f"Subiendo {filename} a Google Drive vía Web App...")
    
    max_retries = 3
    for attempt in range(1, max_retries + 1):
        try:
            with open(file_path, "rb") as f:
                file_bytes = f.read()
                base64_data = base64.b64encode(file_bytes).decode("utf-8")
                
            payload = {
                "filename": filename,
                "base64": base64_data
            }
            
            r = requests.post(f"{web_app_url}?token={token}", json=payload, timeout=45)
            if r.status_code == 200:
                result = r.json()
                if result.get("status") == "success":
                    print(f"¡Éxito! Archivo guardado en Drive con nombre: {result.get('fileName')} (ID: {result.get('fileId')})")
                    return True
                else:
                    print(f"Error al subir (Intento {attempt}/{max_retries}): {result.get('message')}")
            else:
                print(f"Web App devolvió error HTTP {r.status_code} (Intento {attempt}/{max_retries}): {r.text}")
        except Exception as e:
            print(f"Error de red al subir a Google Drive (Intento {attempt}/{max_retries}): {e}")
            
        if attempt < max_retries:
            time.sleep(3 * attempt)
            
    return False

def check_variable_if_unchecked(page, input_id, label_selector):
    chk = page.locator(f"input#{input_id}")
    if chk.count() > 0:
        is_checked = chk.is_checked()
        if not is_checked:
            print(f"Marcando variable {input_id}...")
            page.locator(label_selector).first.click(force=True)

def corregir_desfase_horas(file_path, horas=-6):
    print(f"Corrigiendo desfase de {horas} horas en el archivo {file_path}...")
    try:
        import openpyxl
    except ImportError:
        print("Instalando la librería 'openpyxl'...")
        os.system("pip install openpyxl")
        import openpyxl
        
    try:
        wb = openpyxl.load_workbook(file_path)
        for sheet in wb.worksheets:
            for r in range(1, sheet.max_row + 1):
                for c in range(1, sheet.max_column + 1):
                    val = sheet.cell(row=r, column=c).value
                    if isinstance(val, str):
                        match = re.match(r'^(\d{2}/\d{2}/\d{2,4}) (\d{2}:\d{2}:\d{2})$', val.strip())
                        if match:
                            date_part = match.group(1)
                            year_len = len(date_part.split('/')[-1])
                            fmt = "%d/%m/%y %H:%M:%S" if year_len == 2 else "%d/%m/%Y %H:%M:%S"
                            try:
                                dt = datetime.strptime(val.strip(), fmt)
                                dt_new = dt + timedelta(hours=horas)
                                sheet.cell(row=r, column=c).value = dt_new.strftime(fmt)
                            except Exception:
                                pass
                    elif isinstance(val, datetime):
                        try:
                            sheet.cell(row=r, column=c).value = val + timedelta(hours=horas)
                        except Exception:
                            pass
        wb.save(file_path)
        print("Corrección de desfase de horas aplicada con éxito.")
    except Exception as e:
        print(f"Error al corregir el desfase de horas en Excel: {e}")

def main():
    user, password, web_app_url, token = obtener_credenciales()
    
    # Calcular fecha de ayer en la zona horaria de México (UTC-6)
    mx_tz = timezone(timedelta(hours=-6))
    now_mx = datetime.now(mx_tz)
    yesterday = now_mx - timedelta(days=1)
    
    # Formato Kendo Calendar data-value: YYYY/M/D (con mes 0-indexado)
    kendo_year = yesterday.year
    kendo_month = yesterday.month - 1
    kendo_day = yesterday.day
    kendo_date_val = f"{kendo_year}/{kendo_month}/{kendo_day}"
    
    formatted_date_file = yesterday.strftime("%d-%m-%Y")
    
    print(f"\nFecha a consultar (Ayer en México): {yesterday.strftime('%d/%m/%Y')} (Kendo: {kendo_date_val})")
    
    with sync_playwright() as p:
        print("Iniciando navegador...")
        # Abrimos en modo headless (invisible) para tareas en segundo plano
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(accept_downloads=True, timezone_id="America/Mexico_City", locale="es-MX", viewport={"width": 1280, "height": 800})
        page = context.new_page()
        
        url_login = "https://clientes.ubiqo.net/Publica/Inicio_Sesion.aspx?ReturnUrl=%2fModulos%2fUltimaUbicacion%2fDefaultReact.aspx"
        print(f"Cargando página de login: {url_login}")
        page.goto(url_login)
        
        # 1. Iniciar sesión
        print("Llenando credenciales...")
        page.locator('input[id*="UserName"]').fill(user)
        page.locator('input[id*="Password"]').fill(password)
        
        print("Haciendo clic en 'Iniciar'...")
        btn_iniciar = page.locator('input[type="submit"], button:has-text("Iniciar"), input[value="Iniciar"]').first
        btn_iniciar.click()
        
        # Esperamos el menú de navegación para asegurar el inicio de sesión
        page.locator('a#btnNavegarReportes, span#reportes').first.wait_for(state="visible", timeout=30000)
        print("Sesión iniciada con éxito.")
        
        # 2. Entrar a Reportes (superior izquierda)
        print("Navegando a 'Reportes'...")
        page.locator('a#btnNavegarReportes, span#reportes').first.click()
        page.locator('div#btnReporteDinamico').first.wait_for(state="visible", timeout=20000)
        time.sleep(2)
        
        # 3. Entrar a Dispositivos (Reporte Dinámico)
        print("Navegando a 'Dispositivos' (Reporte Dinámico)...")
        page.locator('div#btnReporteDinamico').first.click()
        page.locator('div#controlFechaInicio_selectorFechasReporte').first.wait_for(state="visible", timeout=25000)
        time.sleep(3)
        
        # 4. Configurar fechas de calendarios (Inicio y Fin = ayer) utilizando la API de Kendo UI
        print("Configurando fechas en calendarios de Kendo mediante API de Kendo...")
        page.evaluate(f"""() => {{
            var calIni = $("#controlFechaInicio_selectorFechasReporte").data("kendoCalendar");
            if (calIni) {{
                calIni.value(new Date({kendo_year}, {kendo_month}, {kendo_day}));
                calIni.trigger("change");
            }}
            var calFin = $("#controlFechaFin_selectorFechasReporte").data("kendoCalendar");
            if (calFin) {{
                calFin.value(new Date({kendo_year}, {kendo_month}, {kendo_day}));
                calFin.trigger("change");
            }}
        }}""")
        time.sleep(2)
            
        # 5. Seleccionar Grupos -> Unidades Instalaciones
        print("Seleccionando grupo 'Unidades Instalaciones'...")
        # Hacemos clic en la pestaña "Grupos" del árbol
        page.locator('span.rtIn:has-text("Grupos")').first.click()
        
        # Esperamos a que los grupos carguen por AJAX y aparezca el nodo "Unidades Instalaciones"
        print("Esperando a que se carguen los grupos en el árbol...")
        unidades_span = page.locator('span.rtIn:has-text("Unidades Instalaciones")').first
        try:
            unidades_span.wait_for(state="attached", timeout=20000)
        except Exception as e:
            print(f"Advertencia al esperar el nodo del grupo: {e}")
        
        # Seleccionamos la casilla de "Unidades Instalaciones" enviando el evento click al checkbox
        print("Marcando la casilla de 'Unidades Instalaciones'...")
        try:
            checkbox_locator = page.locator('//span[text()="Unidades Instalaciones"]/preceding-sibling::input')
            checkbox_locator.wait_for(state="attached", timeout=15000)
            checkbox_locator.dispatch_event('click')
            print("Casilla 'Unidades Instalaciones' marcada con éxito.")
        except Exception as e:
            print(f"Error al marcar la casilla: {e}")
        time.sleep(3)
        
        # 6. Esperar a que las variables se hagan visibles (luego de seleccionar grupo se cargan por AJAX)
        print("Esperando a que se carguen las variables del reporte...")
        
        # Expandir la sección de Variables (Kendo PanelBar)
        print("Expandiendo sección 'Seleccione las variables que desea consultar'...")
        page.locator('span#encabezadoVariables').first.click(force=True)
        time.sleep(1)
        
        first_label = page.locator("label[for='chk_Altitud_admninistradorDevariablesEquipos']").first
        try:
            first_label.wait_for(state="visible", timeout=20000)
            print("Variables cargadas y listas.")
        except Exception as e:
            print(f"Advertencia: Expiró la espera de carga de variables ({e}). Intentando continuar...")
            
        print("Marcando variables a consultar...")
        check_variable_if_unchecked(page, "chk_Altitud_admninistradorDevariablesEquipos", "label[for='chk_Altitud_admninistradorDevariablesEquipos']")
        check_variable_if_unchecked(page, "chk_Parada_admninistradorDevariablesEquipos", "label[for='chk_Parada_admninistradorDevariablesEquipos']")
        check_variable_if_unchecked(page, "chk_Distancia_admninistradorDevariablesEquipos", "label[for='chk_Distancia_admninistradorDevariablesEquipos']")
        check_variable_if_unchecked(page, "chk_Duracion_admninistradorDevariablesEquipos", "label[for='chk_Duracion_admninistradorDevariablesEquipos']")
                    
        # 7. Seleccionar visualización "Día y Ruta"
        print("Seleccionando manera de visualización: 'Día y Ruta'...")
        
        # Expandir la sección de Visualización (Kendo PanelBar)
        print("Expandiendo sección 'Seleccione la manera de visualización'...")
        page.locator('span#contenedorModoVisualizacion').first.click(force=True)
        time.sleep(1)
        
        rad_dia_ruta = page.locator('input#rbtDiaYruta')
        if rad_dia_ruta.count() > 0:
            if not rad_dia_ruta.is_checked():
                try:
                    rad_dia_ruta.click(force=True, timeout=2000)
                except Exception:
                    page.locator('label[for="rbtDiaYruta"]').first.click(force=True)
            
        # 8. Generar Reporte
        print("Haciendo clic en 'Generar Reporte'...")
        page.locator('span#ctl00_cpBody_generarReporte, button.btnGenerarReporteDinamico').first.click()
        
        # 9. Esperar que se generen los resultados (el botón "Exportar Excel" final se hace visible)
        print("Esperando la generación del reporte...")
        btn_excel_final = page.locator('button#btnGenerarExecel')
        try:
            # Esperamos hasta 90 segundos a que aparezca el botón de exportar excel en los resultados
            btn_excel_final.wait_for(state="visible", timeout=90000)
            print("¡Reporte generado con éxito!")
        except Exception:
            print("Advertencia: El tiempo de espera expiró, intentando continuar...")
            
        # Esperar 25 segundos adicionales para asegurar que todos los equipos (8 de 8) se carguen por completo en el grid de Ubiqo
        print("Esperando 25 segundos para permitir la carga completa de todos los equipos en la página...")
        time.sleep(25)
        
        # 10. Exportar Excel
        print("Descargando archivo Excel...")
        try:
            with page.expect_download(timeout=90000) as download_info:
                btn_excel_final.click()
            download = download_info.value
            
            # Nombre de guardado local
            local_filename = f"Reporte Dinamico ({formatted_date_file}).xlsx"
            local_path = os.path.abspath(local_filename)
            download.save_as(local_path)
            print(f"Archivo Excel descargado localmente en: {local_path}")
            
            # Aplicar corrección de desfase de 6 horas si se ejecuta en GitHub Actions (Nube)
            if os.environ.get("GITHUB_ACTIONS") == "true":
                corregir_desfase_horas(local_path, horas=-6)
            
            # 11. Subida a Google Drive vía Web App
            if web_app_url:
                success = subir_a_google_drive(local_path, web_app_url, token)
                if success:
                    try:
                        os.remove(local_path)
                        print("Archivo temporal local eliminado.")
                    except Exception:
                        pass
            else:
                print("\n[ERROR] No se configuró URL del Web App de Apps Script.")
                
        except Exception as e:
            print(f"Error durante la descarga o guardado del reporte: {e}")
            
        print("Cerrando navegador...")
        browser.close()

if __name__ == "__main__":
    main()
